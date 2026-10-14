"""Create the concise meeting brief and reproducibility documentation."""
from pathlib import Path
import json, csv, importlib.metadata
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
S=json.loads((ROOT/'outputs'/'summary.json').read_text())
sources=[
 {'source_id':'S1','publisher':'NYC DOHMH','title':'Restaurant Inspection Results','format':'CSV API extract',
  'url':'https://data.cityofnewyork.us/Health/DOHMH-New-York-City-Restaurant-Inspection-Results/43nn-pn8j/about_data',
  'period':'Record date 2026-10-06','geography':'NYC establishment; borough; ZIP','grain':'One cleaned row per CAMIS',
  'fields_used':'CAMIS, cuisine, borough, ZIP, inspection date, record date, coordinates',
  'use':'Observed restaurant supply','limitation':'Active permits include new applicants and college cafeterias; no revenue, demand or complete closure history'},
 {'source_id':'S2','publisher':'NYC DCWP','title':'Restaurant Delivery App Data','format':'XLSX',
  'url':'https://www.nyc.gov/assets/dca/downloads/xlsx/Restaurant-Delivery-App-Data-Quarterly.xlsx',
  'period':'Q1 2022–Q4 2025','geography':'Citywide covered delivery apps','grain':'Quarter; totals are weekly averages',
  'fields_used':'Deliveries, spending, average order subtotal, merchant platform fees',
  'use':'Delivery volume and per-order fee context','limitation':'Includes some non-restaurant deliveries; no cuisine or borough split; not merchant profit'},
 {'source_id':'S3','publisher':'REBNY','title':'Manhattan Retail Report H1 2026','format':'PDF p. 24',
  'url':'https://www.rebny.com/reports/manhattan-retail-report-first-half-2026/',
  'period':'H1 2023–H1 2026','geography':'16 Manhattan retail corridors','grain':'Corridor × half year',
  'fields_used':'Median annual asking rent per ground-floor square foot',
  'use':'Compare quoted occupancy-cost context','limitation':'General retail asking rents; changing availability sample; not executed restaurant leases'},
 {'source_id':'S4','publisher':'NYC Comptroller','title':'Annual State of the City Economy and Finances 2025','format':'HTML Table 7',
  'url':'https://comptroller.nyc.gov/reports/annual-state-of-the-citys-economy-and-finances-2025/',
  'period':'Years ending August 2019–2025','geography':'NYC restaurants and other eating places','grain':'September–August year',
  'fields_used':'Nominal taxable sales; sales in constant 2019 dollars',
  'use':'Market size and price-adjusted recovery','limitation':'Taxable sales are not total industry revenue; no cuisine, price-tier or individual-restaurant breakdown'},
 {'source_id':'S5','publisher':'NYC Comptroller / Live XYZ','title':'Who Is Minding the Storefronts','format':'HTML narrative',
  'url':'https://comptroller.nyc.gov/reports/whos-minding-the-storefronts/',
  'period':'Status at 2026-04-15; report 2026-06-04','geography':'NYC boroughs','grain':'Borough snapshot',
  'fields_used':'Vacancy percentage by borough; citywide vacancy',
  'use':'Local commercial context','limitation':'All storefronts; not restaurant failure; reported neighborhood rates are not joined to REBNY corridors'},
 {'source_id':'S6','publisher':'Times Square Alliance','title':'August 2026 Monthly Economic Indicator Report','format':'PDF p. 2',
  'url':'https://www.timessquarenyc.org/business-community/market-research-data',
  'period':'August 2026','geography':'Times Square district','grain':'Monthly daily-average pedestrian count',
  'fields_used':'Average daily pedestrian count; rounded year-over-year change',
  'use':'Example of neighborhood foot-traffic context','limitation':'One district; pedestrian counts do not measure restaurant visits or spending'}
]
pd.DataFrame(sources).assign(retrieved_utc_date='2026-10-07').to_csv(ROOT/'source_catalog.csv',index=False)

def p(doc,text,style=None):
    q=doc.add_paragraph(text,style=style)
    q.paragraph_format.keep_with_next=False
    q.paragraph_format.keep_together=False
    q.paragraph_format.widow_control=True
    return q

def link(paragraph,label,url):
    relationship=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),relationship)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'176978');pr.append(color)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);paragraph._p.append(h)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.header_distance=Inches(.22);sec.footer_distance=Inches(.26)
for style in doc.styles:
    if style.type==1:
        style.paragraph_format.keep_with_next=False
        style.paragraph_format.keep_together=False
        style.paragraph_format.page_break_before=False
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    doc.styles[name].font.name='Arial';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.06
doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.underline=False
doc.styles['Subtitle'].font.size=Pt(11);doc.styles['Subtitle'].paragraph_format.space_after=Pt(11)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.size=Pt(12);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(11);doc.styles[name].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(8.5)
doc.core_properties.title='NYC Restaurant Market Platform'
doc.core_properties.subject='Project scope and initial evidence for the internship meeting'
doc.core_properties.author=''
doc.core_properties.keywords='NYC, restaurants, Python, SQL, market research'
foot=sec.footer.paragraphs[0];foot.alignment=2
run=foot.add_run('NYC Restaurant Market Platform   |   ');run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('596670')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
p(doc,'NYC Restaurant Market Platform','Title')
p(doc,'Project scope and initial evidence   •   Prepared 7 October 2026','Subtitle')
p(doc,'Build a public-data platform that helps prospective restaurant owners and existing operators compare possible concepts and locations in New York City. The first version supports exploratory decisions with restaurant supply, retail rent benchmarks and delivery economics. The meeting should confirm this scope and choose the next dashboard implementation.')
p(doc,'User entry points','Heading 1')
t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(2.0);t.columns[1].width=Inches(5.02)
for cell,text in zip(t.rows[0].cells,['What the user knows','What the first version returns']):cell.text=text
for left,right in [
 ('A cuisine or restaurant concept','Compare that cuisine’s recorded supply across boroughs, then inspect ZIP-level counts.'),
 ('A borough or ZIP code','Explore the local cuisine mix and compare relevant cost context where coverage exists.'),
 ('Only that they want to open a restaurant','Start with a citywide overview and narrow the analysis using optional cuisine and location filters.')]:
    cells=t.add_row().cells;cells[0].text=left;cells[1].text=right
for idx,row in enumerate(t.rows):
    for cell in row.cells:
        tcPr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
        for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');marg.append(e)
        tcPr.append(marg)
        if idx==0:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEFF1');tcPr.append(shade)
        for q in cell.paragraphs:
            q.paragraph_format.space_after=Pt(1);q.paragraph_format.line_spacing=1.05
            for run in q.runs:run.font.size=Pt(10);run.bold=idx==0
    if idx==0:
        rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
p(doc,'Research workflow and completed work','Heading 1')
p(doc,'The working assumption is that users currently compare disconnected listings and reports manually; this needs validation with prospective users. The proposed workflow lets them enter a known preference, compare the available evidence, inspect its period and coverage, and export a shortlist for further research.')
p(doc,'The current package reads six public sources, prepares 13 analysis tables, runs SQL aggregations and produces six charts. A local interactive prototype supports all three entry points, CSV downloads and an illustrative rent calculation. The restaurant snapshot contains 31,469 distinct active permit records; 3,860 lack cuisine information. Source data extracts, cleaning code, SQL and quality checks are included.')
p(doc,'Next iteration and acceptance criteria','Heading 1')
p(doc,'Rebuild the core supply view in one agreed BI tool, using the provided CSVs. Validate borough and cuisine filtering against the SQL totals, show missing data, and display the source date and unit beside each chart. Test all three user entry paths. A successful prototype returns consistent comparisons and makes unavailable information explicit. Tableau and Power BI instructions are included; native workbooks have not yet been created.')
p(doc,'Decisions for the meeting','Heading 1')
p(doc,'Confirm the first-version audience and scope, select Tableau or Power BI, and agree whether the next demand measure should be resident population or transit activity. Confirm the expected website format and evaluation criteria. Detailed revenue forecasts, labor-cost estimates and a composite opportunity score require additional evidence and are outside this first iteration.')

doc.add_page_break()
p(doc,'Initial findings','Title')
p(doc,'Delivery growth needs a per-order view','Heading 1')
p(doc,'From Q4 2024 to Q4 2025, average weekly deliveries grew 18.3%. The average subtotal less merchant platform fees rose from $24.08 to $24.26 per delivery, a 0.8% increase. This supports displaying fees alongside volume. The amount is before other restaurant costs and does not establish profit. [S2]')
q=doc.add_paragraph();q.paragraph_format.space_after=Pt(3);q.add_run().add_picture(str(ROOT/'charts'/'delivery_order_economics.png'),width=Inches(6.4))
p(doc,'Location comparisons need rent context','Heading 1')
p(doc,'In H1 2026, Times Square’s median retail asking rent was $1,328 per ground-floor square foot per year, versus $200 in the Financial District, a 6.6-fold difference. These benchmarks support cost comparisons; they do not confirm restaurant suitability or the rent available for an individual space. [S3]')
q=doc.add_paragraph();q.paragraph_format.space_after=Pt(3);q.add_run().add_picture(str(ROOT/'charts'/'rent_selected_corridors.png'),width=Inches(6.4))
q=p(doc,'Sources   ')
q.paragraph_format.space_after=Pt(0)
for i in [1,2,3]:
    source=sources[i];link(q,f'[{source["source_id"]}] {source["publisher"]}',source['url']);q.add_run('   ')
for run in q.runs:run.font.size=Pt(8.5)
p(doc,'Full source URLs, periods, definitions and limits are in source_catalog.csv and DATA_GUIDE.md.','Caption')
for root in [doc.element,doc.styles.element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(ROOT/'NYC_Restaurant_Project_Brief.docx')

readme='''# NYC Restaurant Market Platform

This package prepares the first project discussion with actual public data, a small analysis pipeline and a working local dashboard. The intended users are prospective restaurant owners and existing restaurant operators across NYC.

## Open first

Read `NYC_Restaurant_Project_Brief.docx` for the English meeting brief. Double-click `NYC_Restaurant_Explorer.html` in a browser for the interactive prototype. It works offline because the prepared data are embedded. The source links require an internet connection. `MEETING_NOTES_ZH.md` explains how to present and defend the work.

## What is implemented

The prototype supports cuisine-first, location-first and undecided entry paths. It filters establishment counts by borough, source cuisine and ZIP code. It also compares Manhattan retail asking rents, calculates illustrative base rent for a user-entered floor area, and compares delivery quarters. Citywide sales, storefront vacancy and Times Square pedestrian counts provide context. CSV downloads work locally.

Six public sources have been read. The package preserves the public establishment extract and the selected numerical captures used in the analysis. Python cleans the inputs; SQLite creates six reusable analytical views. The outputs include 13 core CSV tables, six charts in PNG and SVG, a SQLite database and a source catalog. Data quality checks reconcile establishment counts and validate the headline calculations.

## Reproduce the analysis

Python 3.10+ is recommended. Use the dependency versions recorded in `requirements.txt`. The standard library supplies SQLite and HTTP downloads.

```bash
python build_analysis.py
python make_charts.py
python build_dashboard.py
python make_materials.py
```

Run from this directory or use an absolute script path. The scripts find inputs relative to their own location. These commands work offline using the bundled establishment extract and checked numerical captures when full reports are absent. Complete third-party reports are linked rather than redistributed. To deliberately refresh sources, run `python fetch_sources.py` first. The public sources can revise history. A refresh may require schema or period updates; do not mix old outputs with newly downloaded raw data. The current build intentionally checks the reviewed delivery periods and DOHMH snapshot date. Review the frozen narrative and dashboard labels before changing those guards for a newer release.

The dashboard QA script requires Node and Playwright with a Chromium browser, separately from the Python analysis dependencies. Run `node verify_dashboard.cjs` after installing those prerequisites in your own environment. Validation results are in `qa/`; running the QA script also writes screenshots there. Native Tableau and Power BI workbooks are not included; use `BI_BUILD_GUIDE.md` to create one from the prepared tables.

## Package structure

| File or directory | Purpose |
| --- | --- |
| `data/raw/` | Public establishment extract and count; source URLs, timestamps and SHA256 checksums |\n| `data/source_extracts/` | Selected numerical captures from the five reports, with integrity hashes |
| `data/clean/` | Six cleaned source tables |
| `outputs/` | Six SQL views, citywide cuisine counts, summary and validation results |
| `analysis.sql` | Aggregations, denominator definitions and year-over-year calculations |
| `nyc_restaurants.sqlite` | Queryable tables and views |
| `charts/` | Six figures, each in PNG and SVG |
| `source_catalog.csv` | Six sources with coverage, grain, fields, use and limitations |
| `DATA_GUIDE.md` | Definitions, cleaning rules, join restrictions and findings |
| `BI_BUILD_GUIDE.md` | A feasible next-week dashboard implementation |

## Practical limits

This is an exploratory comparison prototype. Current sources do not measure unmet demand, restaurant profit, actual individual lease costs, or the probability a proposed restaurant will succeed. Cuisine counts are supply shares, not revenue market shares. Price-tier revenue, chef and server wages, detailed demographics and a citywide foot-traffic layer have not been added. A hosted website, user accounts and automatic refresh are also future work.
'''
(ROOT/'README.md').write_text(readme)

guide='''# Data guide and initial findings

Prepared 7 October 2026. Use the source catalog and download manifest for exact URLs and source versions. All monetary amounts are USD. Missing numerical values are blank in CSV and null in SQLite. Preserve identifiers such as CAMIS and ZIP as text.

## The evidence and its meaning

The current DOHMH extract has 296,270 inspection/violation rows but 31,469 unique CAMIS establishments. Of these, 31,374 have a recognized borough. Manhattan has 12,565, Brooklyn 7,998, Queens 7,046, the Bronx 2,642 and Staten Island 1,123. The source identifies supply recorded under active permits, including uninspected applicants and college cafeterias. It does not establish current trading, sales or demand. Of all establishments, 3,860 lack cuisine, 494 lack a five-digit ZIP, and 3,858 are marked as never inspected. [S1]

Across covered delivery apps, Q4 2025 weekly delivery volume was 18.3% above Q4 2024. Average subtotal rose from $29.71 to $30.66 per delivery; merchant platform fees rose from $5.63 to $6.40. The difference rose from $24.08 to $24.26, or 0.8%. Merchant fee growth accounts arithmetically for about 80.6% of the subtotal increase per delivery. This is before food, labor, packaging, rent and other costs. Some covered deliveries involve non-restaurant items, and the merchant/order mix can change. [S2]

H1 2026 median asking rents range from $200 in the Financial District to $2,500 on Upper Fifth Avenue among the reported corridors. Times Square is $1,328, 6.64 times the Financial District figure. These are annual asking rents per ground-floor square foot for retail space. They are not signed restaurant leases. The changing available-space sample prevents a same-property interpretation of the time series. [S3]

For September 2024–August 2025, restaurant and other eating-place taxable sales were $28.544 billion nominally and $22.182 billion in constant 2019 dollars. The latter was 3.6% below September 2018–August 2019. This is useful macro context; it cannot be divided into cuisine or restaurant-price-tier performance. [S4]

Storefront vacancy was 11.0% citywide on 15 April 2026. Borough figures were 12.9% in Manhattan, 12.3% in Brooklyn, 9.5% in Queens, 8.0% in the Bronx and 7.5% on Staten Island. These describe all storefronts, so the platform should label them as commercial context rather than restaurant closures. [S5]

The Times Square Alliance reported 240,257 average daily pedestrians in August 2026, with a rounded 2% increase over August 2025. This is a single-district context indicator and does not measure restaurant transactions. [S6]

These findings suggest that the first platform should help users compare evidence, with cost and coverage visible beside supply. A defensible opportunity ranking needs a compatible demand measure and user preferences. Low competition alone is insufficient.

## Data model and joins

| Table or view | Grain | Key and intended use |
| --- | --- | --- |
| restaurants | One CAMIS in the current source snapshot | Unique `camis`; filter by `borough`, `cuisine`, `zipcode` |
| borough_supply | One borough | `borough`; totals include unknown cuisine |
| cuisine_borough | Borough × known cuisine | `borough`, `cuisine`; denominators exclude unknown cuisine |
| zip_cuisine | Borough × ZIP × cuisine | Three-key aggregate; missing ZIP and borough excluded |
| citywide_cuisine CSV | One known cuisine | Counts across all source records, excluding missing cuisine |
| delivery_quarterly / delivery_analysis | One citywide quarter | `year`, `quarter`; do not join to restaurant rows |
| rent_corridor / rent_comparison | Corridor × half year | `corridor`, `year`, `half`; no assumed ZIP mapping |
| citywide_sales / sales_analysis | September–August year | `year_end`; `period_start`, `period_end` make the coverage explicit |
| storefront_vacancy | Borough snapshot | `borough`, `as_of`; label as a separate commercial context measure |
| pedestrian_context | Times Square monthly indicator | `area`, `month`; not a general NYC demand table |

Keep the source fact tables separate. Joining a citywide quarterly amount to every restaurant and summing would multiply it by the number of restaurants. Joining a REBNY corridor to a ZIP by similar names would create unverified geography. Borough vacancy may be displayed beside a borough supply total, but the different dates and populations must remain visible.

## Cleaning and calculation rules

The DOHMH query groups by CAMIS and the current profile fields, keeping the maximum inspection date and record date. This removes repeated violation rows at the source. The build checks the distinct count against a separate source count query and aborts on duplicate CAMIS profiles. It trims text, preserves the publisher's cuisine categories, replaces missing cuisine or unrecognized borough with `Unknown`, and keeps only five-digit formatted ZIPs. This is a formatting check, not geographic validation. The 1900 inspection placeholder becomes a missing date plus `never_inspected=1`. Coordinates are retained with a broad NYC plausibility flag; the flag is not a point-in-boundary check and coordinates are not used for recommendations.

DCWP sheet labels are trimmed and quarter columns are reshaped to one row per quarter. Weekly totals remain weekly averages. Per-delivery values are the agency's reported values, so they need not exactly equal the ratio of two published aggregate averages. The derived after-fee amount is `average_subtotal_usd - merchant_fee_per_delivery_usd`. Year-over-year calculations compare matching quarters using SQL `LAG(...,4)` and guard the denominator.

REBNY medians are parsed from PDF page 24 into long format. The published percentage-change columns are not used. The build recomputes comparisons from the underlying medians and preserves N/A for Harlem and Bleecker Street in H1 2026. The Times Square boundary text and some narrative change statements appear inconsistent in the source, so no exact polygon or narrative-derived change is assigned.

Comptroller sales are extracted from HTML Table 7 and stored in USD millions. The report's constant-dollar series is used directly; the pipeline does not apply another inflation adjustment. Storefront borough percentages are small narrative captures validated against the source passage. Times Square's pedestrian value is extracted from PDF page 2; the reported rounded YoY value is retained separately.

## Field dictionary

`camis` is the establishment identifier; `restaurant_name` is the source DBA. `borough`, `cuisine` and `zipcode` are categorical filters. `latest_inspection_date` is the latest recorded inspection for that establishment, not an opening date. `record_date` is the source snapshot day. `never_inspected` and `coordinate_plausible` are integer flags. `source_id` points to the source catalog.

Counts ending in `_records` or `_count` count establishments, unless a source measure explicitly says deliveries. `share_of_located_records_pct` uses only the five recognized boroughs. `share_of_known_cuisine_records_pct` divides the selected cuisine count by known-cuisine records in that borough. `share_of_cuisine_citywide_pct` divides the borough's cuisine count by that cuisine's count across the five boroughs. All `_pct` fields use 0–100 percentage units; `merchant_fee_share` alone is stored as a 0–1 proportion.

Delivery `weekly_*` fields are average weekly totals within the quarter. `average_*_usd`, `merchant_fee_per_delivery_usd` and `subtotal_less_merchant_fee_usd` are USD per delivery. `prior_year_*` fields are the corresponding values four quarters earlier. All delivery YoY fields are percentage changes from that matching quarter.

`median_asking_rent_usd_sqft_year` and `prior_year_rent` are USD per ground-floor square foot per year. `half` is 1 or 2. Rent `yoy_pct` compares the same half of the previous year and stays missing if either observation is unavailable. The dashboard's monthly illustration is annual asking rate multiplied by entered area and divided by 12.

`nominal_sales_million_usd` is current-dollar taxable sales; `real_sales_2019_million_usd` is the publisher's constant-2019-dollar series. `real_sales_index_2019` sets the year ending August 2019 to 100. `vacancy_pct` is vacant storefronts as a percentage of all storefronts. `average_daily_pedestrians` is the district's published daily average in a month; `reported_yoy_pct_rounded` is a rounded source value, not reconstructed prior-year traffic.

## Source links

'''
for s in sources:guide+=f'[{s["source_id"]}] [{s["publisher"]}: {s["title"]}]({s["url"]}). {s["period"]}.\n\n'
(ROOT/'DATA_GUIDE.md').write_text(guide)

meeting='''# 明天会议怎么使用这份材料

先读两页英文简报，再打开 HTML 原型跑一次演示。你需要能解释项目解决的问题、每个数字代表什么，以及哪些问题还没有数据支持。暂时不用背完整代码。

## 开场可以这样说

I have narrowed the project to a restaurant market exploration platform for both prospective owners and existing operators in New York City. Users can start with a cuisine, a location, or no fixed idea. I have prepared six public sources and a first Python and SQL pipeline, together with a small interactive prototype. The first version focuses on restaurant supply, retail asking-rent context and delivery economics. I would like to use this meeting to confirm the scope, choose the BI tool and agree on the next demand-side dataset.

## 三分钟演示顺序

打开原型，先展示 “I am still exploring”，解释全市概览中的数字是行政记录里的餐饮场所供给。点击 “I have a cuisine”，示例选择 Chinese，然后选 Queens：数量应为 757。点击 “I have a location”，展示 Manhattan 的菜系构成。这样三个不同输入都进入了同一套比较流程。

进入 Delivery trends，保持 Q4 2024 和 Q4 2025，指出配送量增长 18.3%，但平均订单金额扣掉商家平台费后的金额只增加约 0.8%。需要明确，这个金额还没有扣食材、人工、包装和租金，不能称为利润。统计还包括部分非餐饮配送。

进入 Rent comparisons，把面积从 1000 改成 2000 平方英尺，展示成本如何随假设变化。Times Square 与 Financial District 的中位挂牌租金差约 6.6 倍，但这是零售商铺挂牌参考，不代表某一家可用餐馆铺位的实际报价。

## 需要在会议上确认的事

确认第一版是否接受“帮助比较和缩小研究范围”的定位，以及导师希望先看到 Tableau 还是 Power BI。确认网站最终应是公开网页、嵌入 BI 的仪表盘，还是可运行的本地演示。建议下周只新增一种需求侧指标，例如居民人口或交通活动，并先验证与现有地理单元能否对应。还需要确认实习项目最终按分析质量、商业可解释性、网站功能，还是其中某种组合来验收。

## 对方可能问到的问题

如果问为什么不用检查记录的行数数餐馆：同一家店的一次检查可能有多条违规记录。原始来源有 296,270 行，按 CAMIS 归并后是 31,469 个场所。现在的快照只包含活跃状态及部分申请者，所以不能用它重建历年的开关店趋势。

如果问菜系数量能否代表市场份额：目前代表供给数量的占比，不能代表营业收入份额。缺少菜系的信息有 3,860 条，图表保留这一点。更少的同类餐馆也可能意味着需求更少，因此尚未制作“最适合开店”的评分。

如果问为什么没有盈利预测：现有数据缺少单店营收、毛利、具体铺位条款和人工成本。第一版已经有可解释的比较模块，后续可以在取得这些输入后增加经营测算。

如果问你用什么工具：Python 读取 XLSX、HTML、PDF 和 API 数据，完成清洗；SQL 计算汇总、分母和同比；现在用本地 HTML 展示交互。清洗表可以导入 Tableau 或 Power BI，但这两个软件的原生工作簿尚未完成。

## 下周的可执行安排

第一天按会议反馈确定问题和指标，复核原始数据及菜系缺失。第二天在选定的一个 BI 工具中制作餐馆数量与菜系构成视图。第三天加入外卖或租金视图并核对 SQL 总数。第四天测试三类用户入口，补齐来源、时间范围和指标定义。第五天整理反馈和一页书面进度。若需求侧数据当天无法匹配，记录缺口，不延误已验证的供给视图。
'''
(ROOT/'MEETING_NOTES_ZH.md').write_text(meeting)

for package in ['pandas','numpy','openpyxl','lxml','pypdf','matplotlib','python-docx']:
    try:version=importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:continue
    with (ROOT/'requirements.txt').open('a' if package!='pandas' else 'w') as f:f.write(f'{package}=={version}\n')
print('Created brief, source catalog, data guide, meeting notes and README.')
