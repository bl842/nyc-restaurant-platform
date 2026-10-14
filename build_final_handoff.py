"""Build the final BA report and supporting registers from the frozen release."""
from pathlib import Path
import csv,json,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
R=Path(__file__).resolve().parent
D=R/'business_analysis';D.mkdir(exist_ok=True)
def csvout(name,headers,rows):
 with (D/name).open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f);w.writerow(headers);w.writerows(rows)
req=[
('R01','Essential','Compare supply for a known cuisine','Cuisine entry and borough/ZIP filters','S1','Queens + Chinese = 757; five boroughs + Chinese = 2274','Implemented','UI01','US01'),
('R02','Essential','Understand a chosen area','Location entry and cuisine mix','S1','Manhattan all cuisines = 12565; empty selection shows zero and an explanation','Implemented','UI02','US02'),
('R03','Essential','Begin research without fixed preferences','Exploration entry resets supply filters','S1','Reset shows 31374 records across five recognized boroughs','Implemented','UI03','US03'),
('R04','Essential','Compare occupancy cost assumptions','Corridor rent table and floor-area calculator','S3','1000 sq ft at FiDi rate = 16666.67 USD/month; zero area rejected; missing rent stays unavailable','Implemented','UI04','US04'),
('R05','Essential','Assess delivery channel context','Quarter comparison of volume subtotal and merchant fees','S2','Q4 2025 vs Q4 2024 = 18.2798% volume and 0.7633% subtotal less fees growth; same quarter = 0%','Implemented','UI05','US05'),
('R06','Essential','Understand evidence reliability','Source links units periods and coverage notes','S1-S6','Every module exposes period unit and coverage; no profit or success score','Implemented','DOC01','US06'),
('R07','Essential','Reuse and audit results','CSV exports plus reproducible Python/SQL pipeline','S1-S7','Export matches selected ZIP; 17 core tables rebuild; CAMIS unique','Implemented','DATA01/UI06','US07'),
('R08','Recommended','Place results in market context','Taxable sales vacancy and pedestrian context','S4-S6','Real sales comparison = -3.6403%; vacancy labelled all storefronts; pedestrian area labelled Times Square','Implemented','UI07','US08'),
('R09','Essential','Use the prototype on common screen sizes','Responsive layout and labelled controls','None','No horizontal page overflow at 390 and 320 px; no JavaScript errors','Implemented; current browser layout and accessibility user testing pending','UI08','US09'),
('R10','Recommended','Reduce manual BI implementation effort','Native Tableau or Power BI version','Prepared CSVs','Chosen native workbook opens and passes the same count and period controls','Build guide delivered; workbook deferred','Future BI acceptance','US10'),
('R11','Optional','Access a continuously updated shared service','Hosting accounts and scheduled refresh','Publisher access policies','Deployment monitoring and refresh failure tests pass after platform selection','Deferred','Future deployment acceptance','US11'),
('R12','Optional','Receive evidence based opportunity ranking','Demand data and preference based ranking','Additional demand and cost sources','Compatible geography; documented weights; validation against agreed outcomes','Deferred; unsupported by present data','Future model acceptance','US12')]
req.extend([
('R13','Essential','Locate cuisine supply across NYC neighborhoods','NTA choropleth with count/share toggle, borough filter, selection, zoom and export','S1/S7','262 official areas; mapped total 30655; Flushing-Willets Point Chinese count 266 and known denominator 509; zero denominators unavailable','Implemented; DOM and calculation checks; browser layout pending','MAP01','US13'),
('R14','Essential','Inspect where rent evidence applies','Rent hotspot layer over NTA boundaries','S3/S7','16 approximate Manhattan anchors; 14 priced, 2 unavailable; other boroughs have no imputed rent','Implemented; DOM and calculation checks; browser layout pending','MAP02','US14')])
csvout('requirements_traceability.csv' ,['requirement_id','priority','business_need','feature','source','acceptance_criterion','status','verification','user_story_id'],req)
stories=[(r[8],r[0],'As a prospective owner or expansion analyst, I need to '+r[2].lower()+', so I can select the next research question using visible evidence.') for r in req]
csvout('user_stories.csv',['story_id','requirement_id','story'],stories)
kpis=[
('K01','Decision usefulness','Users producing a justified next research question / users completing a scenario','Proposed >=80% in 5-person pilot','Not measured','Moderator scores output against source coverage and caveats','Sponsor/product lead'),
('K02','Task completion','Users completing a scenario without moderator help / users attempting it','Proposed >=80% per scenario','Not measured','Observe cuisine location and undecided tasks separately','BA/researcher'),
('K03','Research time','Median elapsed minutes from entry to exported comparison plus research question','Proposed <=5 minutes; compare with manual baseline','Not measured','Same tasks and time rule for manual and prototype conditions','BA/researcher'),
('K04','Data integrity','Passed deterministic build checks / executed build checks','100% required','12/12 on final rerun','outputs/validation_checks.csv','Data analyst'),
('K05','Traceable requirements','Essential requirements with feature and acceptance criterion / all essential requirements','100% required','10/10 specified; user acceptance pending','requirements_traceability.csv','BA'),
('K06','Evidence coverage','Modules showing source period units and scope / analytical modules','5/5 required','5/5 in prototype content review','Supply maps rent delivery market-context review','Data analyst')]
csvout('kpi_register.csv',['kpi_id','name','definition','target','current_status','measurement_method','proposed_owner'],kpis)
risks=[
('A01','Assumption','Users benefit from area/concept comparisons before a feasibility study','Unvalidated','Run pilot with owners and expansion analysts','Sponsor'),
('A02','Assumption','Current manual research requires reconciling separate reports','Unvalidated workflow hypothesis','Observe a user completing one manual research task','BA'),
('D01','Decision','Accept exploratory comparison scope without opportunity ranking','Pending stakeholder confirmation','Confirm at final review','Sponsor'),
('D02','Decision','Choose Tableau or Power BI and website hosting expectation','Pending stakeholder confirmation','Agree whether offline prototype satisfies current milestone','Sponsor'),
('D03','Decision','Choose next demand source and geographic unit','Pending stakeholder confirmation','Assess resident population versus transit activity and join coverage','Data lead'),
('R01','Risk','Permit counts may be mistaken for open restaurants or sales share','Mitigated in labels; residual interpretation risk','Keep source population and denominator notes beside supply','Data analyst'),
('R02','Risk','Incompatible dates and geography create false local comparisons','Mitigated through separate modules','Do not join corridor rents or citywide delivery to individual restaurants','Data analyst'),
('R03','Risk','Small or changing samples bias asking-rent comparisons','Residual source risk','Retain missing values and ask for actual property quotes','User / analyst'),
('R04','Risk','Frozen values are used as live current estimates','Mitigated through date labels','Refresh in a separate release; revalidate narrative and regression checks','Data maintainer'),
('R05','Risk','Mixed app orders are interpreted as restaurant profitability','Mitigated through metric definitions','Retain non-restaurant coverage note and exclude other costs explicitly','Data analyst'),
('R06','Dependency','Native BI desktop software and an agreed publishing environment','Not provisioned','Use prepared CSVs and build guide after tool selection','Implementation lead')]
csvout('assumptions_risks_decisions.csv',['id','type','item','status','response','proposed_owner'],risks)
# Diagram: branching user entry paths converge on comparison and a human decision.
fig,ax=plt.subplots(figsize=(10,5.0));ax.set_xlim(0,10);ax.set_ylim(0,5);ax.axis('off')
def box(x,y,w,h,t,c='#edf5f5'):
 ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06',facecolor=c,edgecolor='#57717e',linewidth=1));ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=10,color='#173d5c')
def arrow(a,b):ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',color='#57717e',lw=1.3))
for x,title in [(0.15,'Known cuisine\nChoose cuisine'),(3.45,'Known location\nChoose borough or ZIP'),(6.75,'Undecided\nOpen citywide overview')]:
 box(x,3.95,3.0,.72,title);arrow((x+1.5,3.93),(5,3.34))
box(2.0,2.65,6,.7,'Compare neighborhood supply and cuisine mix');arrow((5,2.63),(5,2.2))
box(1.25,1.48,7.5,.7,'Review rent or delivery context\nCheck coverage, dates and missing values');arrow((5,1.46),(5,1.04))
box(2.0,.3,6,.72,'Export comparison and identify the next research question');
fig.tight_layout();fig.savefig(D/'user_workflow.png',dpi=180,bbox_inches='tight');fig.savefig(D/'user_workflow.svg',bbox_inches='tight');plt.close(fig)
# Word report. Explicit page breaks avoid heading pagination markers.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
for st in doc.styles:
 if st.type==1:
  st.paragraph_format.keep_with_next=False;st.paragraph_format.keep_together=False;st.paragraph_format.page_break_before=False
for name,size in [('Normal',10.5),('Title',24),('Heading 1',17),('Heading 2',12),('Caption',9)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
md=[]
def p(t,style=None):
 doc.add_paragraph(t,style);md.append(('## ' if style=='Heading 1' else '### ' if style=='Heading 2' else '# ' if style=='Title' else '')+t+'\n')
def page(title):doc.add_page_break();p(title,'Title')
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for j,h in enumerate(headers):t.rows[0].cells[j].text=h
 for row in rows:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=str(v)
 for ri,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=1
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFEAF0' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.0
    for run in par.runs:run.font.size=Pt(9);run.bold=ri==0
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows)+'\n')
def pic(path,w):doc.add_picture(str(path),width=Inches(w));md.append('![Figure]('+str(path.relative_to(R))+')\n')
p('NYC Restaurant Market Platform','Title')
p('Final business analysis and prototype handoff','Heading 1')
p('Bochao Liang | AI Business Analyst Intern | 13 October 2026')
p('This project provides an evidence-based starting point for prospective restaurant owners and existing operators evaluating NYC concepts or locations. The delivered prototype supports three entry paths and compares restaurant supply, retail asking-rent context and delivery economics. It helps users decide what to investigate next. Commercial feasibility and restaurant success remain unvalidated.')
p('Business problem and objective','Heading 2')
p('Useful evidence is distributed across inspection records and market reports with different dates, units and geographic coverage. Combining those figures without checking their definitions can produce misleading conclusions. The service makes comparable measures visible, adds neighborhood maps with explicit coverage, and lets users export an auditable comparison.')
p('Delivered scope','Heading 2')
p('The package contains seven documented public sources, a reproducible Python and SQLite pipeline, 17 core CSV tables, eight analytical charts, an offline interactive HTML prototype and a business requirements package. The latter includes user journeys, requirements linked to acceptance criteria, KPI definitions, a risk and decision register, and a final review procedure.')
p('Initial recommendation','Heading 2')
p('Use this release as an exploratory comparison platform. Delivery volume grew 18.3% between Q4 2024 and Q4 2025, while order subtotal less merchant platform fees grew only 0.8% per delivery [S2]. Rent benchmarks differ substantially across Manhattan corridors [S3]. These findings support showing costs beside activity and supply rather than treating a single high-growth or low-competition measure as an investment recommendation.')
p('Release boundary','Heading 2')
p('This release reuses the reviewed 7 October market captures and adds official NTA boundaries retrieved on 13 October. The DOHMH record date is 6 October 2026. The prototype is a local HTML file. A native Tableau or Power BI workbook, hosted service and recommendation model are future implementation items. Stakeholder approval and end-user acceptance have not been recorded.')
page('Users and decision workflow')
p('Target users','Heading 1')
table(['Starting point','Decision supported','Minimum input'],[
('Concept known','Compare where that cuisine is recorded and narrow the next research area','Cuisine; borough or ZIP optional'),
('Location known','Understand the area\'s recorded cuisine mix and compare alternatives','Borough or ZIP; cuisine optional'),
('No fixed preference','Explore the citywide supply pattern and form an initial question','No input required')],[1.45,3.65,1.9])
p('An expansion analyst can use the same paths for an existing restaurant business. Chain-specific sales, catchment areas and site suitability require additional company or property data. The three profiles describe intended users; interviews have not yet validated them.')
p('Current workflow hypothesis','Heading 2')
p('A user searches several public reports, copies figures into a spreadsheet, checks units and coverage manually, then tries to compare locations or concepts. This workflow is a hypothesis for user research. No measured baseline or interview evidence is claimed.')
p('Proposed workflow','Heading 2');pic(D/'user_workflow.png',6.65)
p('Cuisine-first and location-first buttons load illustrative defaults, Chinese and Manhattan, which users can change. Entry buttons open the neighborhood map. Separate map filters control spatial supply; the supply table keeps its own borough and ZIP filters. Rent corridors and citywide delivery trends remain independent modules. The final action is an exported comparison and a human research decision; the prototype does not automatically recommend a location.')
page('Business requirements and priorities')
p('Needs translated into behavior','Heading 1')
p('Business needs describe decisions users must make. Features implement those needs. The traceability register links each need to a feature, source, acceptance criterion and verification reference. Priority is a proposed release decision, subject to sponsor confirmation.')
table(['ID','Priority','Business need','Implemented behavior'],[
('R01–03','Essential','Start from a concept, area or no preference','Three entry paths with supply filters and cuisine mix'),
('R04','Essential','Compare occupancy assumptions','Rent comparison and floor-area calculator'),
('R05','Essential','Understand delivery economics','Quarter comparison including merchant fees'),
('R06','Essential','Judge evidence reliability','Source, period, unit and coverage notes'),
('R07','Essential','Audit and reuse comparisons','CSV exports and Python/SQL reproduction'),
('R09','Essential','Read results on common devices','Responsive layout and labelled controls'),
('R08','Recommended','Interpret wider conditions','Sales, vacancy and pedestrian context'),
('R10','Recommended','Continue in an agreed BI tool','Import and validation guide delivered'),
('R13–14','Essential','Locate supply and rent evidence','NTA heat map and corridor hotspots'),
('R11–12','Optional','Share a live service or rank opportunities','Deferred pending hosting and demand evidence')],[.65,1.0,2.5,2.85])
p('Functional and quality boundaries','Heading 2')
p('Unknown categories and unavailable values remain visible. Establishment counts use unique CAMIS identifiers. Per-delivery averages are compared at the quarter grain. Rent estimates require a positive floor area and report base rent only. The prototype runs without an account, API key or network call after opening the local HTML file; external source links require connectivity.')
p('Detailed specifications','Heading 2')
p('business_analysis/requirements_traceability.csv contains 14 requirements and measurable acceptance criteria. business_analysis/user_stories.csv gives the corresponding decision-oriented stories. Essential coverage is 10 of 10 specified requirements; technical checks are separate from user acceptance. The BI guide provides the same control totals for a future native workbook.')
page('Evidence and recommendations')
p('Delivery growth does not establish stronger unit economics','Heading 1')
p('From Q4 2024 to Q4 2025, covered weekly deliveries increased 18.3%. Average subtotal less merchant platform fees changed from $24.08 to $24.26 per delivery, up 0.8%. About 80.6% of the subtotal increase was offset arithmetically by higher merchant fees. Display these measures together. This amount excludes food, labor, packaging, rent and other costs; some app orders are non-restaurant deliveries [S2].')
pic(R/'charts'/'delivery_order_economics.png',4.6)
p('Rent context changes the questions a user should ask','Heading 1')
p('H1 2026 median retail asking rent was $1,328 per ground-floor square foot per year in Times Square and $200 in the Financial District, a 6.64-fold difference. At 1,000 square feet, the illustrative monthly base rents are $110,667 and $16,667. Users still need actual quotes and restaurant-use suitability checks; these observations describe available retail-space samples [S3].')
pic(R/'charts'/'rent_selected_corridors.png',4.6)
p('Supply and market growth require careful interpretation','Heading 2')
p('The snapshot has 31,469 unique establishment IDs, including 95 with unknown borough [S1]. Recorded supply can support comparisons, but does not reveal unmet demand. Real taxable sales in the year ending August 2025 remained 3.6% below the corresponding 2019 period [S4]. The next analytical extension should add a compatible demand measure before developing opportunity scores.')
page('Data sources and quality')
p('Source inventory','Heading 1')
p('The full URLs and field definitions are in source_catalog.csv and DATA_GUIDE.md. Selected numerical captures and the public establishment extract are bundled for offline reproduction. The reports themselves remain linked to their publishers.')
table(['ID and source','Observation coverage','Use and principal limitation'],[
('S1  NYC DOHMH','Record date 6 Oct 2026; NYC establishments','Supply by cuisine/borough/ZIP. Active permits include applicants; no sales or full closure history.'),
('S2  NYC DCWP','Q1 2022–Q4 2025; covered apps citywide','Volume and order economics. Totals are weekly averages; some non-restaurant orders.'),
('S3  REBNY','H1 2023–H1 2026; 16 Manhattan corridors','Median annual ground-floor retail asking rents. Changing sample; no executed restaurant leases.'),
('S4  NYC Comptroller','Years ending August 2019–2025; NYC','Nominal and real taxable sales. No cuisine or price-tier split.'),
('S5  Comptroller / Live XYZ','15 Apr 2026; borough snapshots','All-storefront vacancy. Not restaurant failure or closure rates.'),
('S6  Times Square Alliance','August 2026; Times Square','Average daily pedestrians. One district; no restaurant transactions.'),
('S7  NYC DCP','2020 NTA geography; retrieved 13 Oct 2026','262 statistical areas. Approximate neighborhoods; includes parks and airports.')],[1.45,2.05,3.5])
p('Cleaning and quality controls','Heading 2')
p('Inspection and violation rows are aggregated to one CAMIS profile; conflicting duplicate profiles stop the build. Identifiers remain text. Missing cuisine and borough become Unknown; invalid-format ZIPs remain missing. The 1900 inspection placeholder becomes a missing date with a never-inspected flag. Snapshot integrity hashes protect the small report captures.')
p('Of 31,469 records, 3,860 lack cuisine, 494 lack a five-digit ZIP and 3,858 are marked never inspected. These categories overlap and must not be summed. The five recognized boroughs contain 31,374 records. Quality findings are retained instead of silently removing incomplete establishments.')
p('Model and refresh rules','Heading 2')
p('The six market source facts retain their original grains. Restaurant coordinates are matched to full-resolution NTA polygons; 30,655 records match, while 814 remain excluded with reasons. Citywide quarterly delivery figures are not allocated to individual restaurants. Corridor rent data have no validated ZIP mapping. A refresh is a controlled new release: download, inspect schema and periods, rerun cleaning and SQL, reconcile totals, review narratives and rebuild the prototype before publishing.')
page('KPIs and acceptance plan')
p('Success measures','Heading 1')
p('Technical correctness is tested in this release. Business usefulness requires observation of intended users. Pilot targets below are proposed acceptance thresholds, not achieved results or evidence of commercial impact.')
table(['Measure','Definition or target','Current evidence'],[
('Data integrity','All deterministic build checks pass','12/12 passed on final rerun'),
('Requirement traceability','All essential needs have behavior and criteria','10/10 specified'),
('Evidence coverage','All five modules show source, period, units and scope','5/5 present in content review'),
('Task completion','Proposed: at least 4 of 5 users complete each scenario unaided','Not measured'),
('Decision usefulness','Proposed: at least 4 of 5 users state a justified next question','Not measured'),
('Research time','Proposed: median <=5 minutes per task; compare with manual baseline','Not measured')],[1.55,3.35,2.1])
p('Technical acceptance controls','Heading 2')
p('Control totals include 31,374 recognized-borough records, 12,565 Manhattan records and 757 Queens Chinese-cuisine records. Rent controls include $16,666.67 monthly at 1,000 square feet using the Financial District rate, rejection of zero area and unavailable current medians for two corridors. Delivery controls include 18.2798% volume growth, 0.7633% subtotal-less-fee growth and zero change for identical quarters.')
p('User acceptance scenarios','Heading 2')
p('Ask a cuisine-first participant to compare Chinese-cuisine supply across two boroughs and export one ZIP comparison. Ask a location-first participant to describe Manhattan cuisine mix and explain why a citywide delivery value cannot be interpreted as Manhattan performance. Ask an undecided participant to explore the overview and identify one next research question supported by a source and limitation.')
p('Record completion without help, elapsed time, errors and the final explanation. A response passes usefulness review only when it names a source-backed observation, an appropriate limitation and a concrete next investigation. Do not record a user as passing merely for clicking through. Store actual observations in the provided blank UAT log; technical test output does not substitute for these observations.')
p('Release evidence','Heading 2')
p('outputs/validation_checks.csv records build controls. qa/final_verification.json records the final technical verification status and timestamps. qa/ui_checks.json retains the prior passed browser result. Browser QA remains unavailable for the map release. New map calculation and DOM tests are recorded separately; the prior browser pass does not validate the modified page. business_analysis/kpi_register.csv defines the measurements, proposed owners and unmeasured business outcomes.')
page('Risks decisions and handoff')
p('Assumptions and dependencies','Heading 1')
p('The audience definitions and manual workflow remain hypotheses. The existing sources can support exploratory comparisons, while demand, property suitability, cuisine-level revenues and store-level margins are unavailable. A native BI build depends on tool selection and desktop access. Production hosting requires a deployment environment and an agreed refresh owner.')
table(['Risk','Response'],[
('Counts interpreted as sales or open businesses','Label permit population and distinct-CAMIS denominator; keep missing categories visible.'),
('Geography and periods appear interchangeable','Keep source facts separate; show period and geography beside each module.'),
('Asking rent interpreted as an obtainable lease','Show units and sample limits; obtain individual property quotes before feasibility work.'),
('Frozen data interpreted as live','Display original observation dates; require a reviewed new release after refresh.'),
('Delivery difference interpreted as profit','Name the metric precisely and state excluded costs and mixed-order coverage.')],[2.35,4.65])
p('Decisions requested at final review','Heading 2')
p('Confirm the exploratory first-release scope and intended users. Decide whether the local prototype meets the current milestone or whether a hosted service is required. Choose Tableau or Power BI for a native implementation. Agree on a first demand-side measure and geographic unit. Approve or revise the proposed user acceptance thresholds. The decision register records these items as pending; no sponsor approval is assumed.')
p('Handoff and review sequence','Heading 2')
p('Open this report, then demonstrate the three entry paths and the two map layers in NYC_Restaurant_Explorer.html. Show one supply export, one rent-area calculation and one delivery-period comparison. Review the requirements register against the demonstration and inspect the validation evidence. Record sponsor decisions and user-test observations before changing the release status to accepted.')
p('Maintenance responsibilities','Heading 2')
p('Proposed responsibilities: the sponsor owns scope and acceptance; the BA owns user needs and traceability; the data analyst owns source definitions and reconciliation; the implementation lead owns BI or hosting. These are role assignments for agreement, not confirmed staffing. The package includes code, data definitions, rebuild commands and change notes so another analyst can reproduce and extend it.')
page('Neighborhood maps and spatial evidence')
p('Cuisine concentration at a useful local scale','Heading 1')
p('The map covers 262 official 2020 NTAs, including 197 residential statistical areas. Full-resolution point-in-polygon matching assigns 30,655 of 31,469 establishment records, or 97.4%. The 814 exclusions comprise 95 unknown boroughs, 707 other records with unusable coordinates, 11 outside boundaries and one borough mismatch. These exclusion reasons are mutually exclusive [S1, S7].')
pic(R/'charts'/'cuisine_neighborhood_map.png',4.8)
p('Flushing-Willets Point has 266 mapped Chinese-cuisine records among 509 records with known cuisine, a local supply share of 52.3%. Chinatown-Two Bridges has 112 such records, and Sunset Park Central has 108. These counts describe recorded supply; they do not measure demand or prove market saturation.')
p('Rent geography and interaction','Heading 2')
p('The rent layer displays 16 approximate Manhattan corridor anchors with 14 current medians and two unavailable values. It keeps neighborhood polygons neutral because corridor benchmarks cannot establish an NTA-wide rent level. Other boroughs have no rent observations in this source. Users can switch layers, filter boroughs, select areas, zoom, pan and export map data. Detailed methods and acceptance controls are in MAP_GUIDE.md and qa/map_checks.json.')

for root in [doc.element,doc.styles.element]:
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.core_properties.title='NYC Restaurant Market Platform Final Handoff';doc.core_properties.author='Bochao Liang'
doc.save(R/'NYC_Restaurant_Final_Handoff.docx');(R/'FINAL_HANDOFF.md').write_text('\n'.join(md),encoding='utf-8')
csvout('user_acceptance_log.csv',['participant_code','scenario','date','completed_unaided','elapsed_minutes','source_correct','limitation_correct','next_research_question','issues','reviewer','decision'],[])
print('Final report and BA registers generated')
